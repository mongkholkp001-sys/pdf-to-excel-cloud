import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def export_records_to_excel(records, output_path):
    wb = openpyxl.Workbook()
    
    # Sheet 1: Standard
    ws1 = wb.active
    ws1.title = 'ข้อมูลแบบมาตรฐาน'
    
    # Sheet 2: Detailed
    ws2 = wb.create_sheet(title='ข้อมูลแบบแยกละเอียด')
    
    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    data_font = Font(name='Segoe UI', size=10)
    thin_side = Side(border_style='thin', color='D9D9D9')
    data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    # Setup Sheet 1
    headers1 = [
        'ลำดับ', 'เลขประจำตัวประชาชน', 'รหัสประจำบ้าน', 'ชื่อ-นามสกุล', 'เพศ',
        'สัญชาติ', 'วันเดือนปีเกิด', 'อายุ', 'สถานภาพ', 'มารดา', 'เลขบัตรมารดา',
        'สัญชาติมารดา', 'บิดา', 'เลขบัตรบิดา', 'สัญชาติบิดา', 'ที่อยู่ตามทะเบียนบ้าน',
        'สำนักทะเบียน', 'วันที่แจ้ง/ย้ายเข้า', 'สำนักทะเบียนที่ออก', 'วันที่พิมพ์'
    ]
    ws1.append(headers1)
    
    for r in records:
        full_name = f"{r.get('title','')}{r.get('fname','')} {r.get('lname','')}".strip()
        if not full_name:
            full_name = r.get('full_name', '')
            
        ws1.append([
            r.get('no', 1),
            r.get('cid', ''),
            r.get('hid', ''),
            full_name,
            r.get('gender', ''),
            r.get('nat', 'ไทย'),
            r.get('dob', ''),
            r.get('age', ''),
            r.get('status', ''),
            r.get('mother_name', ''),
            r.get('mother_cid', '-'),
            r.get('mother_nat', 'ไทย'),
            r.get('father_name', ''),
            r.get('father_cid', '-'),
            r.get('father_nat', 'ไทย'),
            r.get('address', ''),
            r.get('reg_office', ''),
            r.get('move_date', ''),
            r.get('issue_office', ''),
            r.get('print_date', '')
        ])
        
    # Setup Sheet 2
    headers2 = [
        'ลำดับ', 'เลขประจำตัวประชาชน', 'รหัสประจำบ้าน', 'คำนำหน้า', 'ชื่อตัว', 'นามสกุล',
        'เพศ', 'สัญชาติ', 'วันเดือนปีเกิด', 'อายุ', 'สถานภาพ',
        'มารดา', 'เลขบัตรมารดา', 'สัญชาติมารดา',
        'บิดา', 'เลขบัตรบิดา', 'สัญชาติบิดา',
        'บ้านเลขที่', 'หมู่ที่', 'ซอย', 'ถนน', 'ตำบล/แขวง', 'อำเภอ/เขต', 'จังหวัด',
        'สำนักทะเบียน', 'วันที่แจ้ง/ย้ายเข้า', 'สำนักทะเบียนที่ออก', 'วันที่พิมพ์'
    ]
    ws2.append(headers2)
    
    for r in records:
        ws2.append([
            r.get('no', 1),
            r.get('cid', ''),
            r.get('hid', ''),
            r.get('title', ''),
            r.get('fname', ''),
            r.get('lname', ''),
            r.get('gender', ''),
            r.get('nat', 'ไทย'),
            r.get('dob', ''),
            r.get('age', ''),
            r.get('status', ''),
            r.get('mother_name', ''),
            r.get('mother_cid', '-'),
            r.get('mother_nat', 'ไทย'),
            r.get('father_name', ''),
            r.get('father_cid', '-'),
            r.get('father_nat', 'ไทย'),
            r.get('addr_no', ''),
            r.get('moo', '-'),
            r.get('soi', '-'),
            r.get('road', '-'),
            r.get('subdistrict', ''),
            r.get('district', ''),
            r.get('province', ''),
            r.get('reg_office', ''),
            r.get('move_date', ''),
            r.get('issue_office', ''),
            r.get('print_date', '')
        ])
        
    for ws in [ws1, ws2]:
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            ws.row_dimensions[row[0].row].height = 22
            for cell in row:
                cell.font = data_font
                cell.border = data_border
                val_str = str(cell.value or '')
                if cell.column == 1 or len(val_str) <= 6 or cell.column in [5, 6, 8, 9, 11, 12, 14, 15]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len * 2.0, 12)
            
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path


def export_land_records_to_excel(records, output_path):
    """
    Exports Land Department asset investigation records in the exact 25-column format requested.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    headers = [
        'ลำดับ',
        'พิมพ์เมื่อ',
        'เลขบัตรประชาชน',
        'รายชื่อ',
        'ผลการสืบทรัพย์',
        'ประเภทหลักประกัน',
        'เลขที่โฉนด',
        'หน้าสำรวจ',
        'เลขที่ดิน',
        'ตำบล/แขวง',
        'อำเภอ/เขต',
        'จังหวัด',
        'อาคารโรงเรือน (ประเภทและเลขที่)',
        'เนื้อที่',
        'ชื่อผู้ถือกรรมสิทธิ์',
        'อายุ',
        'ประเภทการได้มา',
        'ทุนทรัพย์จดทะเบียน (บาท)',
        'ราคาประเมินที่ดิน (บาท)',
        'ราคาประเมินสิ่งปลูกสร้าง (บาท)',
        'ผู้รับจำนอง/ภาระผูกพัน',
        'วงเงินจำนอง (บาท)',
        'สถานะการอายัด',
        'สำนักงานที่ดิน',
        'เบอร์โทรศัพท์'
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    data_font = Font(name='Segoe UI', size=10)
    thin_side = Side(border_style='thin', color='D9D9D9')
    data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)

    for r in records:
        ws.append([
            r.get('no', 1),
            r.get('print_date', '-'),
            str(r.get('cid', '-')),
            r.get('full_name', '-'),
            r.get('search_result', 'พบทรัพย์'),
            r.get('asset_type', '-'),
            str(r.get('deed_no', '-')),
            str(r.get('survey_page', '-')),
            str(r.get('land_no', '-')),
            r.get('tambon', '-'),
            r.get('amphoe', '-'),
            r.get('province', '-'),
            r.get('building_info', '-'),
            r.get('area_text', '-'),
            r.get('owner_formatted', '-'),
            str(r.get('age', '-')),
            r.get('acquisition_type', '-'),
            r.get('reg_capital', 0),
            r.get('land_value', 0),
            r.get('building_value', 0),
            r.get('mortgage_info', '-'),
            r.get('mortgage_amount', '-'),
            r.get('seizure_status', 'ไม่มี'),
            r.get('land_office', '-'),
            str(r.get('phone', '-'))
        ])

    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        # Calculate line count for multiline cells
        max_lines = 1
        for c in row:
            if c.value and '\n' in str(c.value):
                max_lines = max(max_lines, str(c.value).count('\n') + 1)
        ws.row_dimensions[row[0].row].height = max(22, max_lines * 18)

        for cell in row:
            cell.font = data_font
            cell.border = data_border
            col_idx = cell.column
            if col_idx in [1, 2, 3, 5, 7, 8, 9, 14, 16, 17, 25]:
                cell.alignment = center_align
                if col_idx in [3, 25]:
                    cell.number_format = '@'
            elif col_idx in [18, 19, 20]:
                cell.alignment = right_align
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'
            else:
                cell.alignment = left_align

    for col in ws.columns:
        col_idx = col[0].column
        col_letter = get_column_letter(col_idx)
        max_line_len = 10
        for cell in col:
            val_str = str(cell.value or '')
            for line in val_str.splitlines():
                max_line_len = max(max_line_len, len(line))
        ws.column_dimensions[col_letter].width = min(max(max_line_len * 1.4, 13), 55)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
